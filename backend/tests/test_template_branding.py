import base64
import zipfile
from io import BytesIO

import pytest
from docx import Document
from docx.shared import Inches

from app.modules.templates.branding import BrandingError, apply_template_branding, header_summary

PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg=="
)


def _template_docx(with_logo: bool = True) -> bytes:
    document = Document()
    header = document.sections[0].header
    header.paragraphs[0].text = "I.E. República del Perú · UGEL 03 · Año 2026"
    if with_logo:
        header.add_paragraph().add_run().add_picture(BytesIO(PNG_1X1), width=Inches(0.4))
    document.sections[0].footer.paragraphs[0].text = "Formando ciudadanos con valores"
    document.add_paragraph("Contenido del formato institucional")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _generated_docx() -> bytes:
    document = Document()
    document.sections[0].header.paragraphs[0].text = "Cabecera Avendia"
    document.sections[0].footer.paragraphs[0].text = "Página 1 de 1"
    document.add_heading("Sesión de aprendizaje", level=1)
    document.add_paragraph("Cuerpo generado por Avendia con sus estilos.")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def test_branding_copies_header_footer_and_logo_keeping_the_body():
    merged = apply_template_branding(_template_docx(), _generated_docx())

    with zipfile.ZipFile(BytesIO(merged)) as archive:
        assert archive.testzip() is None
        names = set(archive.namelist())
        assert "word/header_institucional.xml" in names
        assert "word/footer_institucional.xml" in names
        assert any(name.startswith("word/media/institucional_header_") for name in names)
        content_types = archive.read("[Content_Types].xml").decode()
        assert "/word/header_institucional.xml" in content_types
        assert 'Extension="png"' in content_types

    result = Document(BytesIO(merged))
    assert "I.E. República del Perú" in result.sections[0].header.paragraphs[0].text
    assert result.sections[0].footer.paragraphs[0].text == "Formando ciudadanos con valores"
    assert result.paragraphs[0].text == "Sesión de aprendizaje"
    assert result.paragraphs[1].text == "Cuerpo generado por Avendia con sus estilos."


def test_branding_keeps_generated_footer_when_template_has_none():
    document = Document()
    document.sections[0].header.paragraphs[0].text = "Solo cabecera"
    document.add_paragraph("x")
    output = BytesIO()
    document.save(output)
    merged = apply_template_branding(output.getvalue(), _generated_docx())
    result = Document(BytesIO(merged))
    assert result.sections[0].header.paragraphs[0].text == "Solo cabecera"
    assert result.sections[0].footer.paragraphs[0].text == "Página 1 de 1"


def test_branding_rejects_invalid_files():
    with pytest.raises(BrandingError):
        apply_template_branding(b"not a zip", _generated_docx())


def test_header_summary_reports_text_and_logo():
    summary = header_summary(_template_docx())
    assert summary["has_logo"] is True
    assert "República del Perú" in str(summary["header_text"])
    assert "Formando ciudadanos" in str(summary["footer_text"])
